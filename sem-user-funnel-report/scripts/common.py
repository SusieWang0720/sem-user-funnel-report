# -*- coding: utf-8 -*-
"""共用工具：广告系列内容解析 + 报表样式"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

COLUMNS = ['投放渠道', 'UIN', '首次访问时间', '广告系列名称', '广告系列内容',
           '区域', '关键词类', '细分词类', '广告系列字词', '浏览器语言',
           '注册用户-神策', '创建应用', '消耗用户数-神策', '付费用户-神策']

WIDTHS = [10, 14, 12, 14, 34, 16, 10, 14, 26, 10, 13, 10, 15, 13]

DATE_COL_IDX = 2  # 首次访问时间在第3列(0-based 2)


def parse_content(content):
    """'区域-Search-关键词类-细分-日期' -> (区域, 关键词类, 细分词类)
    用 '-Search-' 切分而非按下标取，兼容 '巴西--Search-' 这类双横线脏数据。"""
    if not isinstance(content, str) or '-Search-' not in content:
        return '', '', ''
    region, rest = content.split('-Search-', 1)
    parts = rest.split('-')
    return region.rstrip('-'), parts[0], parts[1] if len(parts) > 1 else ''


def apply_parse(df):
    """重新解析拆分列（每次修改广告系列内容后都要调用）"""
    parsed = df['广告系列内容'].apply(lambda c: parse_content(c))
    df['区域'] = parsed.apply(lambda t: t[0])
    df['关键词类'] = parsed.apply(lambda t: t[1])
    df['细分词类'] = parsed.apply(lambda t: t[2])
    return df


def style_workbook(path):
    """深色表头 / Arial / 居中 / 冻结首行 / 自动筛选 / 日期 yyyy/m/d"""
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.fill = PatternFill('solid', fgColor='1A1A1A')
            c.font = Font(name='Arial', bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name='Arial')
                c.alignment = Alignment(horizontal='center')
            row[DATE_COL_IDX].number_format = 'yyyy/m/d'
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = 'A2'
        for i, wd in enumerate(WIDTHS, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = wd
    wb.save(path)
